"""
PROGRAMADOR DE PRODUCCIÓN - Tejido de punto
===========================================

Qué hace este archivo (en palabras sencillas):
  1. Lee el archivo "datos_entrada.xlsx" (pedidos que salen de SAP B1,
     rutas de cada artículo y lista de máquinas).
  2. Ordena los pedidos: primero los de mayor prioridad y, entre ellos,
     los que se entregan antes.
  3. Pasa cada pedido por sus procesos en orden (ej.: Tejido -> Tintorería
     -> Acabado) y en cada proceso lo pone en la máquina que lo termine
     más pronto.
  4. Guarda el resultado en "programa_produccion.xlsx" con colores,
     alertas de pedidos atrasados y un diagrama de Gantt.

Cómo se usa:
  python programador.py                      (usa datos_entrada.xlsx)
  python programador.py mi_archivo.xlsx      (usa otro archivo)
"""

import sys
from datetime import datetime, timedelta, time

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter


ARCHIVO_ENTRADA = "datos_entrada.xlsx"
ARCHIVO_SALIDA = "programa_produccion.xlsx"


# ---------------------------------------------------------------------------
# 1. LECTURA DE DATOS
# ---------------------------------------------------------------------------

def es_si(valor):
    """Convierte 'Si', 'SÍ', 'x', 1... en Verdadero."""
    return str(valor).strip().lower() in ("si", "sí", "s", "x", "1", "true", "yes")


def leer_datos(ruta):
    hojas = pd.read_excel(ruta, sheet_name=None)
    faltan = {"Pedidos", "Rutas", "Maquinas", "Parametros"} - set(hojas)
    if faltan:
        sys.exit(f"ERROR: al archivo le faltan las hojas: {', '.join(sorted(faltan))}")

    pedidos = hojas["Pedidos"].dropna(subset=["Pedido", "Articulo"])
    rutas = hojas["Rutas"].dropna(subset=["Articulo", "Proceso"])
    maquinas = hojas["Maquinas"].dropna(subset=["Maquina", "Proceso"])

    # Parámetros: hoja con dos columnas "Parametro" y "Valor"
    param = dict(zip(hojas["Parametros"]["Parametro"], hojas["Parametros"]["Valor"]))
    inicio = pd.to_datetime(param.get("Fecha_inicio_programa", datetime.today())).to_pydatetime()
    hora_turno = int(param.get("Hora_inicio_turno", 6))

    # Limpieza de textos para que "tejido " y "Tejido" sean lo mismo
    for df, col in ((rutas, "Proceso"), (maquinas, "Proceso")):
        df[col] = df[col].astype(str).str.strip().str.title()
    for df in (pedidos, rutas):
        df["Articulo"] = df["Articulo"].astype(str).str.strip()

    pedidos["Fecha_entrega"] = pd.to_datetime(pedidos["Fecha_entrega"])
    if "Prioridad" not in pedidos:
        pedidos["Prioridad"] = 3
    pedidos["Prioridad"] = pedidos["Prioridad"].fillna(3)

    return pedidos, rutas, maquinas, inicio, hora_turno


# ---------------------------------------------------------------------------
# 2. CALENDARIO DE CADA MÁQUINA
# ---------------------------------------------------------------------------

class Maquina:
    """Guarda el horario de una máquina y hasta cuándo está ocupada."""

    def __init__(self, fila, inicio_programa, hora_turno):
        self.nombre = str(fila["Maquina"]).strip()
        self.proceso = fila["Proceso"]
        self.horas_dia = min(float(fila.get("Horas_por_dia", 24) or 24), 24)
        self.sabado = es_si(fila.get("Trabaja_sabado", "Si"))
        self.domingo = es_si(fila.get("Trabaja_domingo", "No"))
        self.hora_turno = hora_turno
        disponible = fila.get("Disponible_desde")
        self.libre_desde = max(inicio_programa, pd.to_datetime(disponible).to_pydatetime()) \
            if pd.notna(disponible) else inicio_programa

    def trabaja_el_dia(self, fecha):
        dia = fecha.weekday()  # 0 = lunes ... 5 = sábado, 6 = domingo
        if dia == 5:
            return self.sabado
        if dia == 6:
            return self.domingo
        return True

    def ventana_del_dia(self, fecha):
        """Hora de inicio y fin del turno en un día dado."""
        ini = datetime.combine(fecha.date(), time(self.hora_turno))
        return ini, ini + timedelta(hours=self.horas_dia)

    def sumar_horas_trabajo(self, desde, horas):
        """
        Devuelve (inicio_real, fin) de un trabajo de 'horas' que empieza
        no antes de 'desde', saltando noches y días no laborables.
        """
        actual, inicio_real, restante = desde, None, horas
        # Empezamos revisando el día anterior por si hay un turno nocturno que cruza la medianoche
        dia = datetime.combine(desde.date(), time(0)) - timedelta(days=1)
        while True:
            if self.trabaja_el_dia(dia):
                ini, fin = self.ventana_del_dia(dia)
                if fin > actual:
                    actual = max(actual, ini)
                    if inicio_real is None:
                        inicio_real = actual
                    disponible = (fin - actual).total_seconds() / 3600
                    if restante <= disponible:
                        return inicio_real, actual + timedelta(hours=restante)
                    restante -= disponible
                    actual = fin
            dia += timedelta(days=1)


# ---------------------------------------------------------------------------
# 3. PROGRAMACIÓN (el "cerebro")
# ---------------------------------------------------------------------------

def programar(pedidos, rutas, maquinas_df, inicio, hora_turno):
    maquinas = [Maquina(f, inicio, hora_turno) for _, f in maquinas_df.iterrows()]
    programa, advertencias = [], []

    # Regla: primero prioridad (1 = más urgente), luego fecha de entrega más cercana
    pedidos = pedidos.sort_values(["Prioridad", "Fecha_entrega"])

    for _, ped in pedidos.iterrows():
        pasos = rutas[rutas["Articulo"] == ped["Articulo"]]
        if "Paso" in pasos:
            pasos = pasos.sort_values("Paso")
        if pasos.empty:
            advertencias.append((ped["Pedido"], f"El artículo {ped['Articulo']} no tiene ruta. No se programó."))
            continue

        listo_desde = inicio  # el pedido puede empezar desde el inicio del programa
        for _, paso in pasos.iterrows():
            candidatas = [m for m in maquinas if m.proceso == paso["Proceso"]]
            permitidas = paso.get("Maquinas_permitidas")
            if pd.notna(permitidas) and str(permitidas).strip():
                nombres = {n.strip() for n in str(permitidas).split(";")}
                candidatas = [m for m in candidatas if m.nombre in nombres]
            if not candidatas:
                advertencias.append((ped["Pedido"], f"No hay máquinas para el proceso '{paso['Proceso']}'. "
                                                    f"El pedido quedó incompleto."))
                break

            velocidad = float(paso["Kg_por_hora"])
            alistamiento = float(paso.get("Horas_alistamiento", 0) or 0)
            horas = alistamiento + float(ped["Cantidad_kg"]) / velocidad

            # Probamos todas las máquinas posibles y nos quedamos con la que termina primero
            mejor = None
            for m in candidatas:
                ini, fin = m.sumar_horas_trabajo(max(listo_desde, m.libre_desde), horas)
                if mejor is None or fin < mejor[2]:
                    mejor = (m, ini, fin)
            m, ini, fin = mejor
            m.libre_desde = fin

            programa.append({
                "Pedido": ped["Pedido"],
                "Cliente": ped.get("Cliente", ""),
                "Articulo": ped["Articulo"],
                "Cantidad_kg": ped["Cantidad_kg"],
                "Paso": paso.get("Paso", ""),
                "Proceso": paso["Proceso"],
                "Maquina": m.nombre,
                "Inicio": ini,
                "Fin": fin,
                "Horas": round(horas, 1),
                "Fecha_entrega": ped["Fecha_entrega"],
            })
            espera = float(paso.get("Horas_espera_despues", 0) or 0)
            listo_desde = fin + timedelta(hours=espera)

    return pd.DataFrame(programa), pd.DataFrame(advertencias, columns=["Pedido", "Advertencia"])


def resumen_por_pedido(programa, pedidos):
    filas = []
    for _, ped in pedidos.sort_values(["Prioridad", "Fecha_entrega"]).iterrows():
        ops = programa[programa["Pedido"] == ped["Pedido"]] if not programa.empty else programa
        if ops.empty:
            filas.append({"Pedido": ped["Pedido"], "Cliente": ped.get("Cliente", ""),
                          "Articulo": ped["Articulo"], "Cantidad_kg": ped["Cantidad_kg"],
                          "Fecha_entrega": ped["Fecha_entrega"].date(), "Inicio": None,
                          "Fin_estimado": None, "Dias_holgura": None, "Estado": "SIN PROGRAMAR"})
            continue
        fin = ops["Fin"].max()
        # Holgura: días que sobran (positivo) o que faltan (negativo) frente a la entrega
        entrega_fin_dia = ped["Fecha_entrega"] + timedelta(days=1)
        holgura = (entrega_fin_dia - fin).total_seconds() / 86400
        if holgura < 0:
            estado = "ATRASADO"
        elif holgura < 2:
            estado = "EN RIESGO"
        else:
            estado = "A TIEMPO"
        filas.append({"Pedido": ped["Pedido"], "Cliente": ped.get("Cliente", ""),
                      "Articulo": ped["Articulo"], "Cantidad_kg": ped["Cantidad_kg"],
                      "Fecha_entrega": ped["Fecha_entrega"].date(),
                      "Inicio": ops["Inicio"].min(), "Fin_estimado": fin,
                      "Dias_holgura": round(holgura, 1), "Estado": estado})
    return pd.DataFrame(filas)


def carga_por_maquina(programa, maquinas_df):
    if programa.empty:
        return pd.DataFrame()
    carga = programa.groupby(["Proceso", "Maquina"]).agg(
        Pedidos=("Pedido", "count"), Kg=("Cantidad_kg", "sum"),
        Horas_programadas=("Horas", "sum"), Ocupada_hasta=("Fin", "max")).reset_index()
    return carga.sort_values(["Proceso", "Maquina"])


# ---------------------------------------------------------------------------
# 4. GUARDAR EN EXCEL CON FORMATO
# ---------------------------------------------------------------------------

COLORES_ESTADO = {"ATRASADO": "F8696B", "EN RIESGO": "FFEB84", "A TIEMPO": "63BE7B", "SIN PROGRAMAR": "BFBFBF"}
COLORES_PROCESO = ["4F81BD", "F79646", "9BBB59", "8064A2", "4BACC6", "C0504D", "2C4D75", "D9A441"]


def ajustar_hoja(ws):
    encabezado = PatternFill("solid", fgColor="1F3864")
    for celda in ws[1]:
        celda.fill = encabezado
        celda.font = Font(color="FFFFFF", bold=True)
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col in ws.columns:
        largo = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(largo + 2, 10), 45)
    for fila in ws.iter_rows(min_row=2):
        for c in fila:
            if isinstance(c.value, datetime):
                c.number_format = "dd/mm/yyyy hh:mm"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def dibujar_gantt(wb, programa, inicio):
    """Hoja con una fila por máquina y una columna por día, coloreada por proceso."""
    ws = wb.create_sheet("Gantt")
    if programa.empty:
        return
    dia0 = inicio.date()
    ultimo = programa["Fin"].max().date()
    dias = [dia0 + timedelta(days=i) for i in range((ultimo - dia0).days + 1)]
    procesos = list(dict.fromkeys(programa["Proceso"]))
    color = {p: COLORES_PROCESO[i % len(COLORES_PROCESO)] for i, p in enumerate(procesos)}

    ws.cell(1, 1, "Máquina")
    for j, d in enumerate(dias, start=2):
        c = ws.cell(1, j, d.strftime("%d/%m"))
        ws.column_dimensions[get_column_letter(j)].width = 7
    ws.column_dimensions["A"].width = 18

    for i, (maq, ops) in enumerate(programa.sort_values(["Proceso", "Maquina"]).groupby(
            ["Proceso", "Maquina"], sort=False), start=2):
        ws.cell(i, 1, f"{maq[1]} ({maq[0]})")
        for _, op in ops.iterrows():
            for j, d in enumerate(dias, start=2):
                if op["Inicio"].date() <= d <= op["Fin"].date():
                    c = ws.cell(i, j)
                    c.fill = PatternFill("solid", fgColor=color[op["Proceso"]])
                    c.value = f"{c.value}, {op['Pedido']}" if c.value else str(op["Pedido"])
                    c.font = Font(color="FFFFFF", size=8)

    for celda in ws[1]:
        celda.font = Font(bold=True)
        celda.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "B2"

    fila = ws.max_row + 2
    ws.cell(fila, 1, "Convenciones:").font = Font(bold=True)
    for k, p in enumerate(procesos, start=1):
        ws.cell(fila + k, 1, p).fill = PatternFill("solid", fgColor=color[p])
        ws.cell(fila + k, 1).font = Font(color="FFFFFF", bold=True)


def guardar(ruta, resumen, programa, carga, advertencias, inicio):
    with pd.ExcelWriter(ruta, engine="openpyxl") as xl:
        resumen.to_excel(xl, sheet_name="Resumen_pedidos", index=False)
        orden = ["Proceso", "Maquina", "Inicio"]
        (programa.sort_values(orden) if not programa.empty else programa).to_excel(
            xl, sheet_name="Programa_por_maquina", index=False)
        carga.to_excel(xl, sheet_name="Carga_maquinas", index=False)
        advertencias.to_excel(xl, sheet_name="Advertencias", index=False)

    wb = load_workbook(ruta)
    for ws in wb.worksheets:
        ajustar_hoja(ws)

    ws = wb["Resumen_pedidos"]
    col_estado = [c.value for c in ws[1]].index("Estado") + 1
    for fila in range(2, ws.max_row + 1):
        c = ws.cell(fila, col_estado)
        if c.value in COLORES_ESTADO:
            c.fill = PatternFill("solid", fgColor=COLORES_ESTADO[c.value])
            c.font = Font(bold=True)

    dibujar_gantt(wb, programa, inicio)
    wb.save(ruta)


# ---------------------------------------------------------------------------
# 5. PROGRAMA PRINCIPAL
# ---------------------------------------------------------------------------

def main():
    entrada = sys.argv[1] if len(sys.argv) > 1 else ARCHIVO_ENTRADA
    print(f"Leyendo {entrada} ...")
    pedidos, rutas, maquinas, inicio, hora_turno = leer_datos(entrada)
    print(f"  {len(pedidos)} pedidos, {rutas['Articulo'].nunique()} artículos con ruta, {len(maquinas)} máquinas")

    programa, advertencias = programar(pedidos, rutas, maquinas, inicio, hora_turno)
    resumen = resumen_por_pedido(programa, pedidos)
    carga = carga_por_maquina(programa, maquinas)

    try:
        guardar(ARCHIVO_SALIDA, resumen, programa, carga, advertencias, inicio)
    except PermissionError:
        sys.exit(f"ERROR: cierre el archivo {ARCHIVO_SALIDA} en Excel y vuelva a intentarlo.")

    print("\nResultado:")
    for estado, n in resumen["Estado"].value_counts().items():
        print(f"  {estado:<14} {n} pedidos")
    if not advertencias.empty:
        print(f"  ¡Atención! Hay {len(advertencias)} advertencias (ver hoja 'Advertencias').")
    print(f"\nListo. Abra el archivo: {ARCHIVO_SALIDA}")


if __name__ == "__main__":
    main()
